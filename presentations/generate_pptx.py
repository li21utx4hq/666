from pptx import Presentation
from pptx.util import Inches, Pt
import json
import os

# This script generates a PPTX based on slides_content.json.
# Requirements: pip install python-pptx

OUTPUT_DIR = "presentations_output"
OUTPUT_FILENAME = "基建山河_解码中国地形_地形篇.pptx"
CONTENT_FILE = os.path.join(os.path.dirname(__file__), 'slides_content.json')

with open(CONTENT_FILE, 'r', encoding='utf-8') as f:
    slides = json.load(f)

prs = Presentation()

# Set default slide width/height if needed

for i, s in enumerate(slides, start=1):
    # Use a blank slide layout
    slide_layout = prs.slide_layouts[6]  # 6 is usually blank
    slide = prs.slides.add_slide(slide_layout)

    # Title
    if s.get('title'):
        left = Inches(0.5)
        top = Inches(0.3)
        width = Inches(9)
        height = Inches(1)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = s['title']
        p.font.size = Pt(32)
        p.font.bold = True

    # Subtitle or main content
    if s.get('subtitle'):
        left = Inches(0.5)
        top = Inches(1.2)
        width = Inches(9)
        height = Inches(1)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = s['subtitle']
        p.font.size = Pt(20)
        p.font.italic = False

    # Bullets
    if s.get('bullets'):
        left = Inches(0.7)
        top = Inches(2.2)
        width = Inches(8.6)
        height = Inches(4.5)
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        for j, b in enumerate(s['bullets']):
            p = tf.add_paragraph() if j>0 else tf.paragraphs[0]
            p.text = b
            p.level = 0
            p.font.size = Pt(18)

    # Notes (placed in slide notes)
    if s.get('notes'):
        notes_slide = slide.notes_slide
        notes_text_frame = notes_slide.notes_text_frame
        notes_text_frame.text = s['notes']

# Ensure output directory exists
os.makedirs(OUTPUT_DIR, exist_ok=True)
output_path = os.path.join(OUTPUT_DIR, OUTPUT_FILENAME)
prs.save(output_path)
print(f"Generated PPTX: {output_path}")
